"""Assemble the thesis chapters into a single Word document.

Run:  python scripts/build_thesis.py

Reads the Markdown chapters in docs/06-thesis/ in order and produces
docs/06-thesis/THESIS.docx with heading styles, tables and a title page.

The chapters remain the source of truth; this script is a build step, not a
place to edit content. Re-run it after changing any chapter.
"""

from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
CHAPTERS_DIR = ROOT / "docs" / "06-thesis"
OUTPUT = CHAPTERS_DIR / "THESIS.docx"

CHAPTER_ORDER = [
    "ch1-introduction.md",
    "ch2-literature-review.md",
    "ch3-methodology.md",
    "ch4-design-and-implementation.md",
    "ch5-empirical-validation.md",
    "ch6-results-discussion-conclusion.md",
]

TITLE = (
    "A Fuzzy Multi-Criteria Decision Model for SME Credit Appraisal "
    "in Development Banking"
)
SUBTITLE = (
    "Formalising Narrative Appraisal under Dual Credit-Risk and "
    "Development-Impact Objectives: Evidence from Sri Lanka"
)


def add_title_page(doc: Document) -> None:
    for _ in range(4):
        doc.add_paragraph()

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(TITLE)
    run.bold = True
    run.font.size = Pt(20)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(SUBTITLE)
    run.italic = True
    run.font.size = Pt(13)

    for _ in range(6):
        doc.add_paragraph()

    for text, size, bold in [
        ("AAV Athukorala", 13, True),
        ("Student ID: 34764", 11, False),
        ("", 11, False),
        ("MSc in Information Technology", 12, False),
        ("NSBM Green University", 12, False),
        ("", 11, False),
        ("Supervisor: Dr. Pabudi Abeyrathne", 11, False),
    ]:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(text)
        run.bold = bold
        run.font.size = Pt(size)

    doc.add_page_break()


def add_integrity_note(doc: Document) -> None:
    """A standing note on how results in this document were produced."""
    h = doc.add_heading("Note on Results and Reproducibility", level=1)
    for run in h.runs:
        run.font.color.rgb = RGBColor(0x1A, 0x1A, 0x1A)

    for text in [
        "Every quantitative result reported in this thesis is produced by code in "
        "the accompanying repository and can be regenerated from the raw data. No "
        "value has been entered by hand.",

        "Where an experiment could not be run, the claim is withdrawn and the gap "
        "recorded as a limitation rather than filled with an estimate. Negative "
        "results are reported as they occurred: these include a claimed "
        "contribution withdrawn after the literature review (§2.3.3), an expert "
        "scorecard that failed to discriminate (§5.6), and a hypothesis about the "
        "mechanism of a data anomaly that was tested and rejected (§5.3.4).",

        "Sections of Chapter 6 that depend on elicitation data are marked as "
        "pending where that data was not obtained. Scores computed under "
        "placeholder criterion weights are labelled as such throughout and are "
        "not reported as findings.",
    ]:
        p = doc.add_paragraph(text)
        p.paragraph_format.space_after = Pt(10)

    doc.add_page_break()


def parse_table(lines: list[str], start: int) -> tuple[list[list[str]], int]:
    """Consume a Markdown table starting at `start`. Returns (rows, next_index)."""
    rows = []
    i = start
    while i < len(lines) and lines[i].strip().startswith("|"):
        raw = lines[i].strip().strip("|")
        cells = [c.strip() for c in raw.split("|")]
        # Skip the |---|---| separator row.
        if not all(re.fullmatch(r":?-{2,}:?", c) for c in cells if c):
            rows.append(cells)
        i += 1
    return rows, i


def add_table(doc: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    width = max(len(r) for r in rows)
    table = doc.add_table(rows=0, cols=width)
    table.style = "Light Grid Accent 1"

    for r_i, row in enumerate(rows):
        cells = table.add_row().cells
        for c_i in range(width):
            text = clean_inline(row[c_i]) if c_i < len(row) else ""
            cells[c_i].text = text
            for p in cells[c_i].paragraphs:
                for run in p.runs:
                    run.font.size = Pt(9)
                    if r_i == 0:
                        run.bold = True


def clean_inline(text: str) -> str:
    """Strip Markdown inline markers that Word will not render."""
    text = re.sub(r"\*\*(.+?)\*\*", r"\1", text)
    text = re.sub(r"(?<!\*)\*(?!\*)(.+?)(?<!\*)\*(?!\*)", r"\1", text)
    text = re.sub(r"`(.+?)`", r"\1", text)
    text = re.sub(r"\[(.+?)\]\((.+?)\)", r"\1", text)
    return text.strip()


def add_rich_paragraph(doc: Document, text: str, style: str | None = None):
    """Add a paragraph, preserving **bold** and `code` runs."""
    p = doc.add_paragraph(style=style)
    for part in re.split(r"(\*\*.+?\*\*|`.+?`)", text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            p.add_run(part[2:-2]).bold = True
        elif part.startswith("`") and part.endswith("`"):
            run = p.add_run(part[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(9.5)
        else:
            p.add_run(re.sub(r"\[(.+?)\]\((.+?)\)", r"\1", part))
    return p


def render_chapter(doc: Document, path: Path) -> None:
    lines = path.read_text(encoding="utf-8").split("\n")
    i = 0
    in_blockquote = False

    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if not stripped:
            i += 1
            in_blockquote = False
            continue

        # Tables
        if stripped.startswith("|"):
            rows, i = parse_table(lines, i)
            add_table(doc, rows)
            doc.add_paragraph()
            continue

        # Headings
        if stripped.startswith("#"):
            level = len(stripped) - len(stripped.lstrip("#"))
            text = clean_inline(stripped.lstrip("#").strip())
            doc.add_heading(text, level=min(level, 4))
            i += 1
            continue

        # Horizontal rule
        if re.fullmatch(r"-{3,}", stripped):
            doc.add_paragraph()
            i += 1
            continue

        # Blockquote - rendered as an indented italic call-out
        if stripped.startswith(">"):
            text = clean_inline(stripped.lstrip(">").strip())
            if text:
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Inches(0.35)
                run = p.add_run(text)
                run.italic = True
                run.font.size = Pt(10)
            in_blockquote = True
            i += 1
            continue

        # Lists
        m = re.match(r"^(\s*)([-*]|\d+\.)\s+(.*)", line)
        if m:
            indent, marker, text = m.groups()
            style = "List Number" if marker[0].isdigit() else "List Bullet"
            p = add_rich_paragraph(doc, text, style=style)
            if len(indent) >= 2:
                p.paragraph_format.left_indent = Inches(0.6)
            i += 1
            continue

        # Body text - join wrapped lines into one paragraph
        buffer = [stripped]
        i += 1
        while i < len(lines):
            nxt = lines[i].strip()
            if (not nxt or nxt.startswith(("#", "|", ">", "-", "*"))
                    or re.match(r"^\d+\.\s", nxt)):
                break
            buffer.append(nxt)
            i += 1

        add_rich_paragraph(doc, " ".join(buffer))

    doc.add_page_break()


def main() -> int:
    missing = [c for c in CHAPTER_ORDER if not (CHAPTERS_DIR / c).exists()]
    if missing:
        print("Missing chapters: " + ", ".join(missing))
        return 1

    doc = Document()

    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    style.paragraph_format.space_after = Pt(8)
    style.paragraph_format.line_spacing = 1.15

    add_title_page(doc)
    add_integrity_note(doc)

    for name in CHAPTER_ORDER:
        path = CHAPTERS_DIR / name
        print(f"  rendering {name}")
        render_chapter(doc, path)

    doc.save(OUTPUT)

    words = sum(
        len((CHAPTERS_DIR / c).read_text(encoding="utf-8").split())
        for c in CHAPTER_ORDER
    )
    print(f"\nWrote {OUTPUT.relative_to(ROOT)}")
    print(f"  {len(CHAPTER_ORDER)} chapters, ~{words:,} words")
    print("\nChapters remain the source of truth. Edit the .md files and re-run.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
