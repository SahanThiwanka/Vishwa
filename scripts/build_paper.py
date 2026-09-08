"""Assemble the conference paper and supervisor briefing into Word documents.

Run:  python scripts/build_paper.py

Reuses the Markdown-to-docx renderer from build_thesis.py so the paper, the
briefing and the thesis are produced by one code path and cannot drift in
formatting.
"""

from __future__ import annotations

import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))

from build_thesis import render_chapter  # noqa: E402

DOCUMENTS = [
    {
        "source": ROOT / "docs" / "07-paper" / "paper-term-contamination.md",
        "output": ROOT / "docs" / "07-paper" / "PAPER.docx",
        "title": "Target Contamination in a Widely Used Small-Business "
                 "Credit Benchmark",
        "subtitle": "Evidence from the SBA National Dataset",
        "byline": "A. A. V. Athukorala · NSBM Green University, Sri Lanka",
    },
    {
        "source": ROOT / "docs" / "SUPERVISOR-BRIEFING.md",
        "output": ROOT / "docs" / "SUPERVISOR-BRIEFING.docx",
        "title": "Supervisor Briefing",
        "subtitle": "A Dual-Objective Decision Support Model for SME Credit "
                    "Appraisal",
        "byline": "AAV Athukorala · MSc IT, NSBM Green University",
    },
]


def build(spec: dict) -> None:
    if not spec["source"].exists():
        print(f"  missing {spec['source'].name} - skipped")
        return

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    style.paragraph_format.space_after = Pt(8)
    style.paragraph_format.line_spacing = 1.15

    for text, size, bold, italic in [
        (spec["title"], 18, True, False),
        (spec["subtitle"], 12, False, True),
        ("", 11, False, False),
        (spec["byline"], 11, False, False),
    ]:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(text)
        run.bold, run.italic = bold, italic
        run.font.size = Pt(size)

    doc.add_paragraph()
    render_chapter(doc, spec["source"])
    doc.save(spec["output"])

    words = len(spec["source"].read_text(encoding="utf-8").split())
    print(f"  {spec['output'].relative_to(ROOT)}  (~{words:,} words)")


def main() -> int:
    print("Building documents:")
    for spec in DOCUMENTS:
        build(spec)
    return 0


if __name__ == "__main__":
    sys.exit(main())
